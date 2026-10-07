"""Build docs/yukthix/tracker/YukthiX-Tracker.xlsx from the design docs, the code and the PRs.

    python docs/yukthix/tracker/build_tracker.py [--code PATH ...]

--code: checkouts to scan for rule IDs and wired screens (default: this repo).
Pass every worktree that holds unmerged work, e.g. --code ../exam-app-org --code ../exam-app-login.
Columns marked (you) are kept when the tracker is rebuilt; everything else is regenerated.
"""
import argparse, json, re, subprocess
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
YX = HERE.parent
REPO = YX.parent.parent
OUT = HERE / 'YukthiX-Tracker.xlsx'

# Roadmap step -> design docs it builds (BACKEND-ROADMAP.md). Edit when the roadmap changes.
STEP_DOCS = {'0': ['P01'], '1': ['P12', 'P04'], '2': ['P01', 'P02', 'P06', 'M01'], '3': ['P14', 'P20'],
             '3b': ['M14', 'M08'], '4': ['M02', 'P03', 'P07'], '5': ['M03', 'P07', 'P08']}
DONE, IN_PROGRESS = {'0', '1', '2'}, set()
RULE_STATUSES = ['Built + tested', 'Built, no test cites it', 'Tested, code not linked', 'UI prototype only',
                 'Verify (step done, not linked)', 'In progress', 'Not started', 'Not needed']
TEST_RESULTS = ['Pass', 'Fail', 'Blocked', 'Not tested']
WORK_STATUSES = ['Open', 'In progress', 'Done', 'Parked', 'Not started', 'Accepted (won\'t do)', 'Moved to a step']
RULE_RE = re.compile(r'YX-[A-Z0-9]+(?:-[A-Z0-9]+)*-\d+[a-z]?')
DEF_RE = re.compile(r'^\|\s*\**(YX-[A-Z0-9]+(?:-[A-Z0-9]+)*-\d+[a-z]?)\**\s*\|(.*?)\|')
GIT_RULE_ERE = r'YX-[A-Z0-9]+(-[A-Z0-9]+)*-[0-9]+[a-z]?'  # git grep -E is POSIX: no \d, no (?:
TEST_RE = re.compile(r'(\.spec\.|\.e2e-spec\.|\.test\.|/test/|__tests__)')

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='1F3A5F')
FILLS = {'green': 'C6EFCE', 'amber': 'FFEB9C', 'red': 'FFC7CE', 'grey': 'E7E6E6', 'blue': 'DDEBF7'}


def sh(args, cwd):
    r = subprocess.run(args, cwd=cwd, capture_output=True, text=True, encoding='utf-8', errors='replace')
    return r.stdout if r.returncode == 0 else ''


def clean(s):
    s = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', s)
    return re.sub(r'\s+', ' ', s.replace('**', '').replace('`', '')).strip()


def steps_for(doc):
    return [s for s, docs in STEP_DOCS.items() if doc in docs] or ['6+']


# ---------- sources ----------
def read_docs():
    docs, rules = [], {}
    for f in sorted((YX / 'design').glob('[MPT][0-9][0-9]-*.md')) + sorted((YX / 'design').glob('APX-*.md')):
        code = f.name.split('-')[0] if not f.name.startswith('APX') else '-'.join(f.name.split('-')[:2])
        text = f.read_text(encoding='utf-8')
        title = next((clean(l[2:]) for l in text.splitlines() if l.startswith('# ')), f.stem)
        status = next((clean(l.split('Status:**', 1)[1]) for l in text.splitlines() if 'Status:**' in l), '')
        wave = next((clean(l.split(':**', 1)[1]) for l in text.splitlines() if l.startswith('> **Build wave')), '')
        n = 0
        for line in text.splitlines():
            m = DEF_RE.match(line)
            if m and m.group(1) not in rules:
                rules[m.group(1)] = {'doc': code, 'text': clean(m.group(2))}
                n += 1
        docs.append({'code': code, 'file': f.name, 'title': title, 'status': status[:300], 'wave': wave[:200], 'rules': n})
    return docs, rules


def scan_code(roots):
    cites = {}  # rule -> {'backend': set, 'test': set, 'ui': set}
    for root in roots:
        out = sh(['git', 'grep', '-n', '-o', '-E', GIT_RULE_ERE, '--', 'apps', 'packages'], root)
        for line in out.splitlines():
            path, _, rule = line.split(':', 2)
            kind = 'test' if TEST_RE.search(path) else 'ui' if path.startswith(('packages/yx-ui', 'apps/web')) else 'backend'
            cites.setdefault(rule, {'backend': set(), 'test': set(), 'ui': set()})[kind].add(f'{root.name}:{path}')
    return cites


def rule_status(rule, doc, c):
    if c and c['backend'] and c['test']: return 'Built + tested'
    if c and c['backend']: return 'Built, no test cites it'
    if c and c['test']: return 'Tested, code not linked'
    st = steps_for(doc)
    if any(s in IN_PROGRESS for s in st): return 'In progress'
    if any(s in DONE for s in st): return 'Verify (step done, not linked)'
    if c and c['ui']: return 'UI prototype only'
    return 'Not started'


def read_screens(root):
    base = root / 'packages/yx-ui/src/screens'
    web = '\n'.join(p.read_text(encoding='utf-8', errors='replace') for p in (root / 'apps/web/app/yx').rglob('*.tsx')) if (root / 'apps/web/app/yx').exists() else ''
    rows = []
    for f in sorted(base.rglob('*.tsx')):
        if re.search(r'(\.stories|\.test|-kit)\.tsx$', f.name) or f.parent.name == '_kit': continue  # kits are shared parts, not screens
        src = f.read_text(encoding='utf-8', errors='replace')
        names = re.findall(r'export (?:default )?(?:function|const) ([A-Z]\w+)', src)
        wired = sorted({n for n in names if re.search(rf'\b{n}\b', web)})
        rows.append({'group': f.relative_to(base).parts[0], 'file': str(f.relative_to(root)).replace('\\', '/'),
                     'components': len(names), 'wired': ', '.join(wired)})
    stories = {}
    for f in base.rglob('*.stories.tsx'):
        g = f.relative_to(base).parts[0]
        stories[g] = stories.get(g, 0) + len(re.findall(r'^export const \w+', f.read_text(encoding='utf-8', errors='replace'), re.M))
    return rows, stories


def read_md_table(path, first_col):
    rows = []
    for line in path.read_text(encoding='utf-8').splitlines():
        cells = [clean(c) for c in line.strip().strip('|').split('|')]
        if line.startswith('|') and cells and cells[0] not in (first_col, '') and not set(cells[0]) <= set('-: '):
            rows.append(cells)
    return rows


def read_prs():
    out = sh(['gh', 'pr', 'list', '--state', 'all', '--limit', '30', '--json',
              'number,title,state,headRefName,baseRefName,url,statusCheckRollup,updatedAt'], REPO)
    prs = []
    for p in json.loads(out or '[]'):
        if not p['headRefName'].startswith('feat/yukthix'): continue
        checks = [c.get('conclusion') or c.get('status') for c in p.get('statusCheckRollup') or []]
        ci = 'failing' if 'FAILURE' in checks else 'passing' if checks and all(c in ('SUCCESS', 'SKIPPED', 'NEUTRAL') for c in checks) else 'running' if checks else '—'
        prs.append([p['number'], p['title'], p['state'], ci, p['headRefName'], p['baseRefName'], p['updatedAt'][:10], p['url']])
    return prs


# ---------- workbook helpers ----------
def kept(path):
    """(sheet, key) -> {column: value} for the (you) columns of the previous tracker."""
    if not path.exists(): return {}
    wb, keep = load_workbook(path), {}
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        mine = [i for i, h in enumerate(hdr) if h and '(you)' in str(h)]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None and mine:
                keep[(ws.title, str(row[0]))] = {hdr[i]: row[i] for i in mine if row[i] not in (None, '')}
    return keep


def sheet(wb, title, headers, rows, widths, keep, wrap=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        prev = keep.get((title, str(r[0])), {})
        ws.append([prev.get(h, v) if '(you)' in h else v for h, v in zip(headers, r)])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical='center')
    for col in wrap:
        for c in ws[get_column_letter(col)][1:]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def dropdown(ws, col, values, n):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{n + 1}')


def colour(ws, rng, mapping):
    for text, fill in mapping.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",{rng.split(":")[0]}))'],
                                                       fill=PatternFill('solid', fgColor=FILLS[fill])))


# ---------- build ----------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--code', action='append', type=Path, default=[])
    roots = [p.resolve() for p in (ap.parse_args().code or [REPO])]
    keep = kept(OUT)
    docs, rules = read_docs()
    cites = scan_code(roots)
    screens, stories = read_screens(roots[0])
    wb = Workbook()
    wb.remove(wb.active)
    today = date.today().isoformat()

    # Rules
    rrows = []
    for rid, r in sorted(rules.items(), key=lambda kv: (kv[1]['doc'], kv[0])):
        c = cites.get(rid)
        auto = rule_status(rid, r['doc'], c)
        ex = sorted((c['backend'] | c['test']) if c else set())[:3]
        rrows.append([rid, r['doc'], rid.rsplit('-', 1)[0], ', '.join(steps_for(r['doc'])), r['text'][:600], auto, None, None,
                      len(c['backend']) if c else 0, len(c['test']) if c else 0, len(c['ui']) if c else 0, '\n'.join(ex), None])
    rh = ['Rule ID', 'Doc', 'Area', 'Roadmap step', 'Rule (from the design doc)', 'Status (auto)', 'Override (you)', 'Status',
          'Code files', 'Test files', 'UI files', 'Where (first 3)', 'Notes (you)']
    ws = sheet(wb, 'Rules', rh, rrows, [16, 8, 12, 9, 80, 24, 22, 24, 8, 8, 8, 50, 30], keep, wrap=(5, 12, 13))
    for i in range(2, len(rrows) + 2):
        ws[f'H{i}'] = f'=IF(G{i}<>"",G{i},F{i})'
    dropdown(ws, 'G', RULE_STATUSES, len(rrows))
    colour(ws, f'F2:H{len(rrows) + 1}', {'Built +': 'green', 'Built,': 'amber', 'Tested,': 'amber', 'Verify': 'red',
                                         'In progress': 'blue', 'Not needed': 'grey'})
    n_rules = len(rrows) + 1

    # Design docs
    drows = [[d['code'], d['title'], ', '.join(steps_for(d['code'])), d['rules'], None, None, None, None, None, d['wave'], d['status'], d['file'], None] for d in docs]
    dh = ['Doc', 'Title', 'Roadmap step', 'Rules', 'Built + tested', 'Built, no test', 'Verify', 'Not started', '% built',
          'Build wave', 'Design status', 'File', 'Notes (you)']
    ws = sheet(wb, 'Design docs', dh, drows, [9, 50, 9, 7, 9, 9, 8, 9, 8, 40, 60, 36, 30], keep, wrap=(2, 10, 11))
    for i in range(2, len(drows) + 2):
        f = lambda s: f'=COUNTIFS(Rules!$B$2:$B${n_rules},$A{i},Rules!$H$2:$H${n_rules},"{s}*")'
        ws[f'E{i}'], ws[f'F{i}'], ws[f'G{i}'], ws[f'H{i}'] = f('Built +'), f('Built,'), f('Verify'), f('Not started')
        ws[f'I{i}'] = f'=IF(D{i}=0,"",E{i}/D{i})'
        ws[f'I{i}'].number_format = '0%'

    # Roadmap
    road = read_md_table(YX / 'BACKEND-ROADMAP.md', 'Step')
    roadrows = []
    for r in road:
        st = r[0]
        docs_in = STEP_DOCS.get(st, [])
        state = 'Done' if st in DONE else 'In progress' if st in IN_PROGRESS else (r[3] if len(r) > 3 and r[3] else 'Not started')
        roadrows.append([st, r[1], r[2] if len(r) > 2 else '', state, None, None, None])
    oh = ['Step', 'What', 'Design docs', 'Status', 'Rules in its docs', 'Built + tested', 'Notes (you)']
    ws = sheet(wb, 'Roadmap', oh, roadrows, [7, 90, 22, 22, 10, 10, 30], keep, wrap=(2, 4, 7))
    for i, r in enumerate(roadrows, 2):
        if r[0] in STEP_DOCS:
            docs_in = STEP_DOCS[r[0]]
            ws[f'E{i}'] = '=' + '+'.join(f'COUNTIF(Rules!$B$2:$B${n_rules},"{d}")' for d in docs_in)
            ws[f'F{i}'] = '=' + '+'.join(f'COUNTIFS(Rules!$B$2:$B${n_rules},"{d}",Rules!$H$2:$H${n_rules},"Built +*")' for d in docs_in)

    # Open work + deferred
    work = json.loads((HERE / 'open-work.json').read_text(encoding='utf-8'))
    wrows = [[w['id'], w['area'], w['item'], w['status'], None, w['next'], None, None] for w in work]
    wh = ['ID', 'Area', 'What', 'Status (start)', 'Status (you)', 'Next action', 'Owner (you)', 'Notes (you)']
    ws = sheet(wb, 'Open work', wh, wrows, [8, 26, 90, 14, 16, 40, 14, 30], keep, wrap=(3, 6, 8))
    dropdown(ws, 'E', WORK_STATUSES, len(wrows))
    defer = json.loads((HERE / 'deferred.json').read_text(encoding='utf-8'))
    frows = [[d['id'], d['step'], d['source'], d['date'], d['item'], 'Open (verify)', None, None, None] for d in defer]
    fh = ['ID', 'Step', 'Build job', 'Date', 'Left out / not finished (as reported by the build job)', 'Status (start)', 'Status (you)', 'Owner (you)', 'Notes (you)']
    ws = sheet(wb, 'Deferred', fh, frows, [8, 6, 22, 11, 100, 13, 16, 14, 30], keep, wrap=(5, 9))
    dropdown(ws, 'G', WORK_STATUSES, len(frows))

    # Go-live
    golive = read_md_table(YX / 'GO-LIVE-CHECKLIST.md', '#')
    grows = [[g[0], g[1], g[2] if len(g) > 2 else '', g[3] if len(g) > 3 else '', 'Open', None, None, None] for g in golive]
    gh = ['#', 'Item', 'What to do', 'Notes', 'Status (start)', 'Status (you)', 'Owner (you)', 'Target date (you)']
    ws = sheet(wb, 'Go-live', gh, grows, [5, 26, 80, 50, 12, 14, 14, 14], keep, wrap=(3, 4))
    dropdown(ws, 'F', WORK_STATUSES, len(grows))

    # Screens
    srows = [[s['file'], s['group'], s['components'], 'Yes' if s['wired'] else 'No', s['wired'], None] for s in screens]
    ws = sheet(wb, 'Screens', ['Screen file', 'Group', 'Components', 'Wired to API', 'Wired components', 'Notes (you)'],
               srows, [60, 14, 11, 12, 50, 30], keep, wrap=(5,))
    colour(ws, f'D2:D{len(srows) + 1}', {'Yes': 'green', 'No': 'amber'})

    # PRs & branches
    prs = read_prs()
    ws = sheet(wb, 'PRs', ['PR', 'Title', 'State', 'CI', 'Branch', 'Base', 'Updated', 'Link', 'Notes (you)'],
               prs, [6, 60, 9, 9, 24, 24, 11, 48, 30], keep, wrap=(2,))
    colour(ws, f'D2:D{len(prs) + 1}', {'passing': 'green', 'failing': 'red', 'running': 'amber'})
    br = [[str(r.name), sh(['git', 'branch', '--show-current'], r).strip(), sh(['git', 'log', '-1', '--format=%h %s (%cs)'], r).strip()] for r in roots]
    sheet(wb, 'Code scanned', ['Checkout', 'Branch', 'Last commit'], br, [24, 28, 90], keep)

    # Backlog (DevOps style: Epic > Feature > Story), Plan and Coverage
    items = [it for f in sorted((HERE / 'stories').glob('*.json')) for it in json.loads(f.read_text(encoding='utf-8'))]
    n_back = 0
    if items:
        byid = {it['id']: it for it in items}
        def chain(it):
            seen, out = set(), []
            while it and it['id'] not in seen:
                seen.add(it['id']); out.append(it); it = byid.get(it.get('parent', ''))
            return out[::-1]
        def key(it):
            return [(c.get('order', 9999), c['id']) for c in chain(it)]
        items.sort(key=key)
        brows = []
        for it in items:
            ch = chain(it)
            epic = next((c['title'] for c in ch if c['type'] == 'Epic'), '')
            feat = next((c['title'] for c in ch if c['type'] == 'Feature'), '')
            pad = {'Epic': '', 'Feature': '    ', 'Story': '        '}.get(it['type'], '')
            brows.append([it['id'], it.get('order'), it['type'], pad + it['title'], epic, feat, it.get('state', 'New'), None, None,
                          it.get('phase', ''), it.get('points'), it.get('priority'), it.get('story', ''),
                          '\n'.join('• ' + a for a in it.get('acceptance', [])), ', '.join(it.get('rules', [])),
                          ', '.join(it.get('docs', [])), ', '.join(it.get('depends', [])), it.get('evidence', ''), None])
        bh = ['ID', 'Order', 'Type', 'Title', 'Epic', 'Feature', 'State (start)', 'State (you)', 'State', 'Phase', 'Points',
              'Priority', 'User story', 'Acceptance criteria', 'Rules', 'Design docs', 'Depends on', 'Evidence', 'Notes (you)']
        ws = sheet(wb, 'Backlog', bh, brows, [11, 7, 8, 55, 30, 30, 10, 12, 10, 18, 7, 7, 60, 80, 30, 14, 16, 40, 30], keep, wrap=(4, 13, 14, 15, 18))
        n_back = len(brows) + 1
        for i, r in enumerate(brows, 2):
            ws[f'I{i}'] = f'=IF(H{i}<>"",H{i},G{i})'
            if r[2] in ('Epic', 'Feature'):
                for c in ws[i]:
                    c.font = Font(bold=True)
                    c.fill = PatternFill('solid', fgColor=FILLS['blue' if r[2] == 'Epic' else 'grey'])
        dropdown(ws, 'H', ['New', 'Active', 'Done', 'Removed'], len(brows))
        colour(ws, f'I2:I{n_back}', {'Done': 'green', 'Active': 'amber'})
        # Plan: phases in build order
        phases = {}
        for r in brows:
            if r[2] == 'Story':
                p_ = phases.setdefault(r[9] or '(no phase)', {'order': r[1] if r[1] is not None else 9999, 'epics': set()})
                p_['order'] = min(p_['order'], r[1] if r[1] is not None else 9999); p_['epics'].add(r[4])
        BANDS = [(100, 'Step 0'), (200, 'Step 1'), (300, 'UI prototype'), (400, 'Step 2'), (500, 'Step 3'), (600, 'Step 3b'),
                 (700, 'Step 4'), (800, 'Step 5'), (1500, 'Step 6+ HR modules'), (2000, 'ATS & assessments'), (10 ** 9, 'Later')]
        band = lambda o: next(n for lim, n in BANDS if o < lim)
        prow = []
        for ph, v in sorted(phases.items(), key=lambda kv: kv[1]['order']):
            prow.append([ph, band(v['order']), v['order'], '; '.join(sorted(v['epics'])), None, None, None, None, None, None])
        ph_h = ['Phase', 'Roadmap step', 'Starts at order', 'Epics', 'Stories', 'Points', 'Done', 'Active', 'New', '% done (points)']
        ws = sheet(wb, 'Plan', ph_h, prow, [30, 18, 10, 60, 9, 9, 8, 8, 8, 12], keep, wrap=(4,))
        rng = lambda col: f"Backlog!${col}$2:${col}${n_back}"
        for i in range(2, len(prow) + 2):
            st = f'{rng("J")},$A{i},{rng("C")},"Story"'
            ws[f'E{i}'] = f'=COUNTIFS({st})'
            ws[f'F{i}'] = f'=SUMIFS({rng("K")},{st})'
            for col, state in (('G', 'Done'), ('H', 'Active'), ('I', 'New')):
                ws[f'{col}{i}'] = f'=COUNTIFS({st},{rng("I")},"{state}")'
            ws[f'J{i}'] = f'=IF(F{i}=0,"",SUMIFS({rng("K")},{st},{rng("I")},"Done")/F{i})'
            ws[f'J{i}'].number_format = '0%'
        # Coverage: every design rule must sit in at least one story
        in_stories = {r for it in items for r in it.get('rules', [])}
        cov = [[rid, r['doc'], r['text'][:300], 'Not in any story'] for rid, r in sorted(rules.items()) if rid not in in_stories]
        cov += [[rid, '', '', 'Story cites a rule the docs do not define'] for rid in sorted(in_stories - set(rules))]
        sheet(wb, 'Coverage', ['Rule ID', 'Doc', 'Rule', 'Problem'], cov, [16, 8, 90, 34], keep, wrap=(3,))

    # Manual tests (manual-tests.json): the founder's screen-by-screen script; placed right after Summary
    mt = json.loads((HERE / 'manual-tests.json').read_text(encoding='utf-8'))
    mrows = [[m['id'], m['area'], m['screen'], m['url'], m['account'], '\n'.join(f'{i}. {s}' for i, s in enumerate(m['steps'], 1)),
              m['expected'] + (f"\n\nNote: {m['note']}" if m.get('note') else ''), ', '.join(m['stories']), None, None] for m in mt]
    mh = ['ID', 'Area', 'Screen', 'URL', 'Account', 'Steps', 'Expected', 'Stories', 'Result (you)', 'Notes (you)']
    ws_mt = sheet(wb, 'Manual tests', mh, mrows, [10, 18, 30, 34, 28, 70, 60, 18, 13, 40], keep, wrap=(3, 5, 6, 7, 8, 10))
    dropdown(ws_mt, 'I', TEST_RESULTS, len(mrows))
    colour(ws_mt, f'I2:I{len(mrows) + 1}', {'Pass': 'green', 'Fail': 'red', 'Blocked': 'amber'})
    n_mt = len(mrows) + 1

    # Summary (first sheet)
    ws = wb.create_sheet('Summary', 0)
    wb.move_sheet(ws_mt, offset=1 - wb.index(ws_mt))
    ws['A1'] = 'YukthiX product tracker'
    ws['A1'].font = Font(bold=True, size=16)
    ws['A2'] = f'Generated {today} from docs/yukthix (design docs, roadmap, go-live) and the code in: ' + ', '.join(r.name for r in roots)
    rows = [('Design rules', None)]
    rows += [(s, f'=COUNTIF(Rules!$H$2:$H${n_rules},"{s}")') for s in RULE_STATUSES]
    rows += [('All rules', f'=COUNTA(Rules!$A$2:$A${n_rules})'),
             ('% built + tested', f'=COUNTIF(Rules!$H$2:$H${n_rules},"Built + tested")/COUNTA(Rules!$A$2:$A${n_rules})'),
             ('', None), ('Work', None),
             ('Open work items not done', f'=COUNTA(\'Open work\'!A2:A{len(wrows) + 1})-COUNTIF(\'Open work\'!E2:E{len(wrows) + 1},"Done")'),
             ('Deferred items still open', f'=COUNTA(Deferred!A2:A{len(frows) + 1})-COUNTIF(Deferred!G2:G{len(frows) + 1},"Done")-COUNTIF(Deferred!G2:G{len(frows) + 1},"Accepted*")-COUNTIF(Deferred!G2:G{len(frows) + 1},"Moved*")'),
             ('Go-live items not done', f'=COUNTA(\'Go-live\'!A2:A{len(grows) + 1})-COUNTIF(\'Go-live\'!F2:F{len(grows) + 1},"Done")'),
             ('', None), ('UI', None),
             ('Screen files', len(srows)), ('Screen files wired to the API', sum(1 for s in srows if s[3] == 'Yes')),
             ('Storybook stories', sum(stories.values())),
             ('', None), ('Manual tests', None),
             ('Test cases', f"=COUNTA('Manual tests'!A2:A{n_mt})"),
             ('Passed', f"=COUNTIF('Manual tests'!I2:I{n_mt},\"Pass\")"),
             ('Failed', f"=COUNTIF('Manual tests'!I2:I{n_mt},\"Fail\")"),
             ('Blocked', f"=COUNTIF('Manual tests'!I2:I{n_mt},\"Blocked\")"),
             ('Not tested', f"=COUNTA('Manual tests'!A2:A{n_mt})-COUNTIF('Manual tests'!I2:I{n_mt},\"Pass\")-COUNTIF('Manual tests'!I2:I{n_mt},\"Fail\")-COUNTIF('Manual tests'!I2:I{n_mt},\"Blocked\")"),
             ('', None), ('Design', None), ('Design docs', len(drows)), ('PRs (YukthiX branches)', len(prs))]
    if n_back:
        b = lambda col: f'Backlog!${col}$2:${col}${n_back}'
        rows += [('', None), ('Backlog', None), ('Epics', f'=COUNTIF({b("C")},"Epic")'), ('Features', f'=COUNTIF({b("C")},"Feature")'),
                 ('Stories', f'=COUNTIF({b("C")},"Story")')]
        rows += [(f'Stories {st}', f'=COUNTIFS({b("C")},"Story",{b("I")},"{st}")') for st in ('Done', 'Active', 'New')]
        rows += [('Story points total', f'=SUMIFS({b("K")},{b("C")},"Story")'),
                 ('% points done', f'=SUMIFS({b("K")},{b("C")},"Story",{b("I")},"Done")/SUMIFS({b("K")},{b("C")},"Story")'),
                 ('Rules not covered by a story', f"=COUNTA(Coverage!A2:A{len(rules) + 50})")]
    for i, (k, v) in enumerate(rows, 4):
        ws[f'A{i}'], ws[f'B{i}'] = k, v
        if v is None and k: ws[f'A{i}'].font = Font(bold=True)
        if k.startswith('%'): ws[f'B{i}'].number_format = '0.0%'
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 34, 14
    r0 = len(rows) + 6
    ws[f'A{r0}'] = 'How to read this'
    ws[f'A{r0}'].font = Font(bold=True)
    notes = [
        'Rules: every YX-... rule in the design docs. "Built + tested" means a code file and a test file cite the rule ID.',
        '"Verify" means its roadmap step is done but no code or test cites the rule: confirm it is built and add the ID to its test, or build it.',
        'Override (you) on the Rules sheet replaces the automatic status (e.g. "Not needed" with a note).',
        'Manual tests: the founder test script (manual-tests.json). Set Result (you) to Pass, Fail, Blocked or Not tested; empty counts as not tested.',
        'Deferred: things each build job reported as left out. Close, move to a step, or accept each one.',
        'Columns marked (you) are kept when the tracker is regenerated. Everything else is rebuilt from the docs and the code.',
        'Rebuild: python docs/yukthix/tracker/build_tracker.py --code <each worktree with unmerged work>',
    ]
    for j, t in enumerate(notes, r0 + 1):
        ws[f'A{j}'] = t
    wb.calculation.fullCalcOnLoad = True  # formulas have no cached values; Excel computes them on open
    wb.save(OUT)
    print(f'{OUT}: {len(rrows)} rules, {len(drows)} docs, {len(frows)} deferred, {len(wrows)} open work, {len(grows)} go-live, {len(srows)} screens, {len(prs)} PRs, {len(mrows)} manual tests')


if __name__ == '__main__':
    main()
